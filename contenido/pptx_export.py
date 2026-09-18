#!/usr/bin/env python3
"""Exporta un carrusel a .pptx para editarlo en Canva.

Por que no alcanza el PDF: Canva lo importa con un parser propio que (a) pinta
los degradados como rectangulos negros, (b) se come los espacios entre palabras
porque Chromium los escribe como corrimientos de posicion y no como el caracter
espacio, y (c) reemplaza la tipografia. El PPTX en cambio entra a Canva como
cajas de texto y formas nativas, con las palabras separadas y editables.

Como funciona: se abre el HTML en el navegador, se mide cada elemento con su
estilo ya calculado, y con esas medidas se arma el PPTX. Nada se escribe dos
veces: el HTML sigue siendo la fuente.

Uso:  python3 contenido/pptx_export.py contenido/carruseles/<carpeta>
"""
import io
import pathlib
import sys

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

PX = 9525                      # 1 px CSS = 9525 EMU (96 dpi)
W, H = 1080, 1350
CHROME = '/opt/pw-browsers/chromium-1194/chrome-linux/chrome'

# Que se exporta y en que orden. El orden importa: lo que va antes queda abajo,
# asi que los rectangulos tienen que ir antes que el texto que llevan encima.
SPEC = [
    ('.dash',            'rect',  False),
    ('.eyebrow > span',  'text',  False),
    ('h1',               'text',  False),
    ('p.sub',            'text',  True),
    ('ol.pasos b',       'text',  False),
    ('ol.pasos .pt',     'text',  False),
    ('ol.pasos .pa',     'text',  True),
    ('.vs > div',        'rect',  False),
    ('.vs em',           'text',  False),
    ('.vs > div > span', 'text',  True),
    ('.cta',             'rect',  False),
    ('.cta > span',      'text',  False),
    ('.shot',            'image', False),
    ('.brand img',       'image', False),
    ('.brand span',      'text',  False),
    ('.swipe',           'text',  False),
    ('.num',             'text',  False),
]

MEDIR = """
(spec) => {
  const px = v => parseFloat(v) || 0;
  const color = (v, fondo) => {
    const m = v.match(/[\\d.]+/g) || [0, 0, 0];
    const a = m.length > 3 ? +m[3] : 1;
    const c = [+m[0], +m[1], +m[2]];
    return c.map((x, i) => Math.round(a * x + (1 - a) * (fondo ? fondo[i] : 255)));
  };
  // El texto se parte en corridas para conservar los colores de adentro del
  // titular, y en parrafos donde el copy puso un <br>.
  const leer = el => {
    const parrafos = [[]];
    (function walk(n, cs) {
      for (const hijo of n.childNodes) {
        if (hijo.nodeType === 3) {
          const t = hijo.textContent.replace(/\\s+/g, ' ');
          if (t.trim() || t === ' ') parrafos[parrafos.length - 1].push(
            {texto: t, color: color(cs.color), bold: +cs.fontWeight >= 600});
        } else if (hijo.tagName === 'BR') {
          parrafos.push([]);
        } else if (hijo.nodeType === 1) {
          if (hijo.tagName.toLowerCase() === 'svg') {
            parrafos[parrafos.length - 1].push(
              {texto: ' \u2192', color: color(cs.color), bold: true});
          } else {
            walk(hijo, getComputedStyle(hijo));
          }
        }
      }
    })(el, getComputedStyle(el));
    return parrafos.filter(p => p.length);
  };

  return [...document.querySelectorAll('.slide')].map(slide => {
    const base = slide.getBoundingClientRect();
    const cs0 = getComputedStyle(slide);
    const fondo = color(cs0.backgroundColor);
    const piezas = [];
    const rel = r => ({x: r.left - base.left, y: r.top - base.top,
                       w: r.width, h: r.height});

    for (const [sel, tipo, wrap] of spec) {
      for (const el of slide.querySelectorAll(sel)) {
        const cs = getComputedStyle(el);
        const caja = rel(el.getBoundingClientRect());
        if (tipo === 'rect') {
          piezas.push({tipo: 'rect', ...caja, fill: color(cs.backgroundColor),
                       radio: px(cs.borderTopLeftRadius)});
        } else if (tipo === 'image') {
          piezas.push({tipo: 'image', ...caja, sel,
                       idx: [...slide.querySelectorAll(sel)].indexOf(el),
                       radio: px(cs.borderTopLeftRadius)});
        } else {
          // el marcador se dibuja aparte: en el PPTX es un rectangulo relleno
          // detras del texto, uno por renglon que ocupa
          for (const marca of el.querySelectorAll('.mark')) {
            const msc = getComputedStyle(marca);
            const alto = 0.36 * px(msc.fontSize);
            for (const r of marca.getClientRects()) {
              const c = rel(r);
              piezas.push({tipo: 'rect', x: c.x, y: c.y + c.h - alto, w: c.w,
                           h: alto, fill: color(msc.boxShadow.match(/rgba?\\([^)]*\\)/)[0], fondo),
                           radio: 0});
            }
          }
          piezas.push({tipo: 'text', ...caja, wrap, parrafos: leer(el),
                       fuente: cs.fontFamily.split(',')[0].replace(/["']/g, ''),
                       size: px(cs.fontSize), interlinea: px(cs.lineHeight),
                       tracking: px(cs.letterSpacing),
                       mayusculas: cs.textTransform === 'uppercase'});
        }
      }
    }
    return {fondo, piezas};
  });
}
"""


def _rgb(c):
    return RGBColor(*c)


def _texto(slide, p):
    caja = slide.shapes.add_textbox(Emu(int(p['x'] * PX)), Emu(int(p['y'] * PX)),
                                    Emu(int((p['w'] + (6 if p['wrap'] else 60)) * PX)),
                                    Emu(int((p['h'] + 8) * PX)))
    tf = caja.text_frame
    tf.word_wrap = p['wrap']
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, parrafo in enumerate(p['parrafos']):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = PP_ALIGN.LEFT
        par.line_spacing = Pt(p['interlinea'] * 0.75)
        for corrida in parrafo:
            r = par.add_run()
            r.text = corrida['texto'].upper() if p['mayusculas'] else corrida['texto']
            # sin xml:space="preserve" un lector estricto se come el espacio del
            # final de una corrida, y "Era el " + "primer" queda "Era elprimer":
            # exactamente el bug que tiene el PDF en Canva
            r._r.find(qn('a:t')).set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
            r.font.size = Pt(p['size'] * 0.75)
            r.font.name = p['fuente']
            r.font.bold = corrida['bold']
            r.font.color.rgb = _rgb(corrida['color'])
            if p['tracking']:
                # python-pptx no expone el espaciado entre letras: va directo al XML,
                # en centesimas de punto
                r.font._rPr.set('spc', str(int(p['tracking'] * 0.75 * 100)))
    return caja


def _rect(slide, p):
    pill = p['radio'] >= min(p['w'], p['h']) / 2
    forma = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if p['radio'] else MSO_SHAPE.RECTANGLE,
        Emu(int(p['x'] * PX)), Emu(int(p['y'] * PX)),
        Emu(int(p['w'] * PX)), Emu(int(p['h'] * PX)))
    if p['radio']:
        forma.adjustments[0] = 0.5 if pill else p['radio'] / min(p['w'], p['h'])
    forma.fill.solid()
    forma.fill.fore_color.rgb = _rgb(p['fill'])
    forma.line.fill.background()
    forma.shadow.inherit = False
    return forma


def exportar(folder):
    from playwright.sync_api import sync_playwright

    d = pathlib.Path(folder).resolve()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W * PX), Emu(H * PX)
    vacia = prs.slide_layouts[6]

    with sync_playwright() as p:
        exe = CHROME if pathlib.Path(CHROME).exists() else None
        b = p.chromium.launch(executable_path=exe, args=['--no-sandbox'])
        pg = b.new_page(viewport={'width': W, 'height': H}, device_scale_factor=2)
        pg.goto((d / 'index.html').as_uri())
        pg.wait_for_timeout(1200)
        datos = pg.evaluate(MEDIR, SPEC)

        for n, info in enumerate(datos):
            slide = prs.slides.add_slide(vacia)
            fondo = slide.background.fill
            fondo.solid()
            fondo.fore_color.rgb = _rgb(info['fondo'])
            for pieza in info['piezas']:
                if pieza['tipo'] == 'rect':
                    _rect(slide, pieza)
                elif pieza['tipo'] == 'text':
                    _texto(slide, pieza)
                else:
                    el = pg.query_selector_all('.slide')[n].query_selector_all(pieza['sel'])[pieza['idx']]
                    slide.shapes.add_picture(
                        io.BytesIO(el.screenshot()),
                        Emu(int(pieza['x'] * PX)), Emu(int(pieza['y'] * PX)),
                        Emu(int(pieza['w'] * PX)), Emu(int(pieza['h'] * PX)))
        b.close()

    salida = d / 'editable-canva.pptx'
    prs.save(str(salida))
    print('%d slides -> %s (%.0f KB)' % (len(datos), salida.name, salida.stat().st_size / 1024))
    return salida


if __name__ == '__main__':
    exportar(sys.argv[1])
